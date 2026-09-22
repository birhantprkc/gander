#!/usr/bin/env python3
"""
Generates every fixture document the test suite reads.

Three tiers share these files: the JVM and Robolectric tests take them as test
resources, the instrumented tests as assets inside the test APK, and the Python
viewer tests straight off disk. One directory, one generator, so a format only
ever needs describing once.

Everything here is invented. Nothing imitates a real organisation, and no
identifier is registry-shaped: see tests/fixtures/README.md for why that rule
exists.

Output is byte-stable, so regenerating an unchanged fixture produces no diff.
PDFs get ReportLab's invariant flag; the OOXML formats are zips, so their
entry timestamps and core properties are normalised by hand afterwards.

Usage:  python3 tests/fixtures/make_fixtures.py
Needs:  reportlab python-docx openpyxl python-pptx pillow cryptography
"""

import hashlib
import heapq
import hmac
import io
import os
import random
import re
import shutil
import struct
import sys
import wave
import zipfile
import zlib
from datetime import datetime
from math import pi, sin
from pathlib import Path

OUT = Path(__file__).parent / "files"

# Every timestamp written into a fixture. Arbitrary, fixed, and in the past.
EPOCH = datetime(2026, 1, 1, 0, 0, 0)
ZIP_DATE = (2026, 1, 1, 0, 0, 0)

# The word test_pdf_search.py counts. It appears exactly three times in
# six-pages.pdf and nowhere else in it, in three different cases.
NEEDLE_LINES = [
    "This tenancy begins on the first of March.",
    "The Tenancy may be ended by either party.",
    "Nothing in this TENANCY limits the above.",
]


def written(path: Path) -> Path:
    print(f"  {path.relative_to(OUT.parent.parent)}  {path.stat().st_size:,} B")
    return path


# ---------------------------------------------------------------------------
# Zip normalisation, for the three OOXML formats
# ---------------------------------------------------------------------------

ISO_EPOCH = EPOCH.strftime("%Y-%m-%dT%H:%M:%SZ")


def normalize_zip(path: Path) -> None:
    """Rewrites a zip with fixed entry timestamps, order and compression.

    python-docx and friends stamp every entry with the time the file was
    written, so an unchanged fixture would still produce a diff on every run.

    docProps/core.xml is rewritten as well. openpyxl in particular sets
    dcterms:modified to the moment of the save, overwriting whatever the
    workbook properties said, so pinning it before the save does nothing.
    """
    with zipfile.ZipFile(path) as z:
        items = sorted((i.filename, z.read(i.filename)) for i in z.infolist())
    items = [
        (name, re.sub(
            rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)",
            rb"\g<1>" + ISO_EPOCH.encode() + rb"\g<2>",
            data,
        ) if name == "docProps/core.xml" else data)
        for name, data in items
    ]
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in items:
            info = zipfile.ZipInfo(name, date_time=ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            z.writestr(info, data)
    tmp.replace(path)


def fix_core_properties(doc) -> None:
    """Pins the created and modified times an OOXML package records."""
    cp = doc.core_properties
    cp.created = EPOCH
    cp.modified = EPOCH
    cp.last_modified_by = "Gander tests"
    cp.author = "Gander tests"
    cp.revision = 1


# ---------------------------------------------------------------------------
# PDFs
# ---------------------------------------------------------------------------

def pdfs() -> None:
    from reportlab.lib.pagesizes import A3, A4, landscape
    from reportlab.lib.pdfencrypt import StandardEncryption
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    def new(path: Path, pagesize=A4, **kw):
        # invariant drops the creation date and the document id, which are the
        # only two things that would otherwise differ between two identical runs
        return canvas.Canvas(str(path), pagesize=pagesize, invariant=1, **kw)

    # six-pages.pdf: the workhorse. Base-14 fonts only, so no embedded face,
    # and exactly three occurrences of the search needle.
    c = new(OUT / "six-pages.pdf")
    for n in range(1, 7):
        c.setFont("Helvetica-Bold", 18)
        c.drawString(72, 760, f"Alder Court, page {n}")
        c.setFont("Helvetica", 11)
        y = 720
        for line in [
            "A short agreement written only so that a test has something to read.",
            "Every name, address and figure in it is invented.",
        ]:
            c.drawString(72, y, line)
            y -= 18
        if n <= len(NEEDLE_LINES):
            c.drawString(72, y - 12, NEEDLE_LINES[n - 1])
        c.showPage()
    c.save()
    written(OUT / "six-pages.pdf")

    # forty-pages.pdf: long enough that the page band holds a fraction of it,
    # which is what the virtualisation and go-to-page tests need.
    c = new(OUT / "forty-pages.pdf")
    for n in range(1, 41):
        c.setFont("Helvetica-Bold", 16)
        c.drawString(72, 760, f"Section {n}")
        c.setFont("Helvetica", 11)
        c.drawString(72, 730, f"This is page {n} of forty.")
        c.showPage()
    c.save()
    written(OUT / "forty-pages.pdf")

    # embedded-font.pdf: carries its own face, so the text layer must name
    # that face rather than a generic. ReportLab ships Vera under a licence
    # that allows redistribution.
    vera = Path(pdfmetrics.__file__).parent.parent / "fonts" / "Vera.ttf"
    pdfmetrics.registerFont(TTFont("Vera", str(vera)))
    c = new(OUT / "embedded-font.pdf")
    c.setFont("Vera", 20)
    c.drawString(72, 700, "Embedded Vera, not a system face.")
    c.showPage()
    c.save()
    written(OUT / "embedded-font.pdf")

    # cjk.pdf: a CID font named but deliberately NOT embedded, so pdf.js can
    # only draw it by loading Adobe's UniGB CMap out of lib/cmaps. Without
    # those tables the text vanishes from the canvas and the text layer both,
    # silently, which is issue #21.
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    c = new(OUT / "cjk.pdf")
    c.setFont("STSong-Light", 22)
    c.drawString(72, 700, "你好世界")  # ni hao shi jie
    c.setFont("Helvetica", 12)
    c.drawString(72, 660, "The line above must render and be selectable.")
    c.showPage()
    c.save()
    written(OUT / "cjk.pdf")

    # mixed-width.pdf: one A4 page then one A3, to pin that both lay out to the
    # same CSS width. That normalisation was once thought to be issue #20; it is
    # not, because a page shown at one width is equally sharp whatever its paper
    # size. The blur was the single rasterisation, and dense-map.pdf tests it.
    c = new(OUT / "mixed-width.pdf", pagesize=A4)
    c.setFont("Helvetica", 24)
    c.drawString(72, 700, "A4 page")
    c.showPage()
    c.setPageSize(A3)
    c.setFont("Helvetica", 24)
    c.drawString(72, 1000, "A3 page")
    c.showPage()
    c.save()
    written(OUT / "mixed-width.pdf")

    # dense-map.pdf: A3 landscape carrying detail into every corner, which is
    # what a tube map or a site plan is and what issue #20 was reported against.
    # The tile tests need a page whose middle is not blank: a fixture with a
    # line of text at the top correlates to nothing once you zoom past it, and
    # a test comparing two blank regions agrees with itself perfectly.
    c = new(OUT / "dense-map.pdf", pagesize=landscape(A3))
    w, h = landscape(A3)
    c.setLineWidth(0.25)
    c.setStrokeColorRGB(0.78, 0.78, 0.84)
    for x in range(0, int(w), 20):
        c.line(x, 0, x, h)
    for y in range(0, int(h), 20):
        c.line(0, y, w, y)
    c.setStrokeColorRGB(0.1, 0.2, 0.7)
    c.setLineWidth(2)
    for i in range(9):
        c.line(60 + i * 130, 60, 60 + i * 130 + 300, h - 60)
    c.setFillColorRGB(0, 0, 0)
    c.setFont("Helvetica", 3.5)
    for x in range(0, int(w), 100):
        for y in range(0, int(h), 60):
            c.drawString(x + 2, y + 2, f"St {x}/{y}")
    c.showPage()
    c.save()
    written(OUT / "dense-map.pdf")

    # ragged-prose.pdf: issue #22. A text layer only covers its glyphs, so the
    # leading between lines, the white beside a short line, the gutter between
    # two columns and the page margins belong to no span at all, and a finger
    # dragging a selection through one of them lands on nothing. padRows() in
    # pdf.html pads the spans until they tile the page, and this is the page
    # shaped to make it work: ragged right ends, one line of three words, a
    # wide blank before a heading, a two-column block with a gutter down the
    # middle, and a row mixing 16pt with 8pt. Every other PDF fixture here is
    # full-width lines at one size, which padRows covers without trying.
    c = new(OUT / "ragged-prose.pdf")
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, 780, "Ragged prose")
    c.setFont("Helvetica", 11)
    y = 748
    for line in [
        "The first line runs the whole width of the measure and then stops here.",
        "A shorter second line.",
        "Three words.",
        "The fourth line is long again, so the ragged edge above it has somewhere",
        "to be, and the space beside it belongs to no span until padRows runs.",
    ]:
        c.drawString(72, y, line)
        y -= 18

    # The wide blank. Displacement is a fraction of a span's top padding, so it
    # is invisible on an ordinary line gap and shows up on the line after this.
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, 540, "After a wide blank")
    c.setFont("Helvetica", 11)
    y = 508
    for line in [
        "This heading sits below 130 points of nothing, which is where the",
        "vertical padding is largest and any error in it is largest too.",
    ]:
        c.drawString(72, y, line)
        y -= 18

    # Two columns, so a gap inside a row has to be covered as well as the two
    # margins. A selection dragged down the left column passes through it.
    y = 430
    for left, right in [
        ("First entry", "Fourteen"),
        ("Second entry", "Twenty one"),
        ("Third entry", "Three"),
        ("Fourth entry", "Eight"),
    ]:
        c.drawString(72, y, left)
        c.drawString(340, y, right)
        y -= 18

    # One row, two type sizes. Padding each item down by the same amount from
    # where it happens to end would leave a sliver under the smaller one.
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, 330, "Large label")
    c.setFont("Helvetica", 8)
    c.drawString(300, 330, "and a caption beside it, set much smaller")

    # Sideways text, which padRows skips a span at a time. The rest of the page
    # still has to come out covered.
    c.saveState()
    c.rotate(90)
    c.setFont("Helvetica", 10)
    c.drawString(200, -560, "Printed sideways in the margin")
    c.restoreState()

    c.setFont("Helvetica", 11)
    c.drawString(72, 120, "A last line, well clear of everything above it.")
    c.showPage()
    c.save()
    written(OUT / "ragged-prose.pdf")

    # colours.pdf: everything night mode has to get right, in known values so a
    # test can assert exact pixels rather than "darker".
    #
    # Page 1 is a document: white paper, near-black text, a saturated heading and
    # a vector block, all of which turn over, plus a small image that must not.
    # Page 2 is a scan: one image covering the whole page, which must turn over
    # despite being an image, because that is what a scanned book is.
    #
    # Flat colours on purpose. A photograph would make the assertions sample
    # noise, and what is being pinned here is which pixels the filter reached.
    from PIL import Image, ImageDraw
    from reportlab.lib.utils import ImageReader

    def block(rgb, size=(64, 64)):
        return ImageReader(Image.new("RGB", size, rgb))

    W, H = 400, 600
    c = new(OUT / "colours.pdf", pagesize=(W, H))
    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, W, H, stroke=0, fill=1)
    # Up in the top-left corner, so the first zoom tile lands on it and the tile
    # path's own coordinate mapping is exercised rather than assumed.
    c.drawImage(block((255, 0, 255)), 20, 420, width=100, height=100)
    c.setFillColorRGB(20 / 255, 20 / 255, 20 / 255)
    c.setFont("Helvetica", 14)
    c.drawString(160, H - 60, "Body text")
    c.setFillColorRGB(0, 119 / 255, 199 / 255)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(160, H - 100, "Coloured heading")
    c.setFillColorRGB(30 / 255, 150 / 255, 60 / 255)
    c.rect(160, H - 200, 80, 60, stroke=0, fill=1)
    c.drawImage(block((200, 30, 30)), 40, 120, width=180, height=180)
    c.showPage()

    scan = Image.new("RGB", (200, 300), (255, 255, 255))
    for x in range(20, 180):
        for y in range(40, 60):
            scan.putpixel((x, y), (20, 20, 20))
    c.drawImage(ImageReader(scan), 0, 0, width=W, height=H)
    c.showPage()

    # Page 3 is the case that decides night mode's image rule, and the one two
    # simpler rules got wrong. All three of these are images as far as the PDF is
    # concerned, and they want three different things:
    #
    #   the chart   an opaque white background: leaving it alone puts a white
    #               rectangle on a dark page, which is what KOReader shipped and
    #               had reported back as their issue #4986
    #   the mono    a grey photograph, so no saturation to give it away; only the
    #     photo     absence of white paper separates it from a document
    #   the colour  the easy case, and the one a size rule got wrong by treating a
    #     photo     full-page photograph as a page
    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, W, H, stroke=0, fill=1)

    chart = Image.new("RGB", (260, 170), (255, 255, 255))
    pen = ImageDraw.Draw(chart)
    pen.line([(30, 140), (250, 140)], fill=(20, 20, 20), width=2)
    pen.line([(30, 10), (30, 140)], fill=(20, 20, 20), width=2)
    pen.line([(30, 120), (90, 60), (150, 95), (210, 30)], fill=(0, 119, 199), width=3)
    c.drawImage(ImageReader(chart), 30, H - 220, width=260, height=170)

    mono = Image.new("RGB", (160, 120))
    rng = random.Random(11)
    for x in range(160):
        for y in range(120):
            v = rng.randint(20, 150)
            mono.putpixel((x, y), (v, v, v))
    c.drawImage(ImageReader(mono), 30, H - 380, width=160, height=120)

    colour = Image.new("RGB", (160, 120))
    for x in range(160):
        for y in range(120):
            colour.putpixel((x, y), (rng.randint(120, 255), rng.randint(20, 90),
                                     rng.randint(20, 90)))
    c.drawImage(ImageReader(colour), 220, H - 380, width=160, height=120)

    # A product shot: a coloured object on a white background, which is what most
    # photographs in a catalogue or a listing actually are. Mostly white, so the
    # "is it mostly paper" half of the rule says document; strongly coloured, so the
    # saturation half says picture. It is the case that stops that half being dropped.
    product = Image.new("RGB", (150, 110), (252, 252, 252))
    shot = ImageDraw.Draw(product)
    shot.ellipse([30, 20, 120, 90], fill=(214, 68, 24))
    shot.ellipse([52, 34, 78, 54], fill=(250, 186, 96))
    c.drawImage(ImageReader(product), 30, H - 520, width=150, height=110)

    # A second figure, over on the right. The one on the left cannot tell whether the
    # page sample is being mapped back to the page at all, because a wrong mapping
    # still lands somewhere near the top left and still reads as paper. This one is
    # far enough across that a wrong mapping reads nothing and leaves it white.
    right = Image.new("RGB", (200, 130), (255, 255, 255))
    pen2 = ImageDraw.Draw(right)
    pen2.rectangle([20, 20, 60, 110], fill=(20, 20, 20))
    pen2.rectangle([80, 55, 120, 110], fill=(20, 20, 20))
    pen2.rectangle([140, 35, 180, 110], fill=(20, 20, 20))
    c.drawImage(ImageReader(right), 210, H - 520, width=170, height=110)
    c.showPage()

    # Page 4: two photographs that overlap, which is how a collage, a watermarked
    # photo or a figure with an inset is put together. Both are pictures and all of
    # both must survive, the overlap included. Clipping every picture out of one
    # even-odd path gets this wrong - canvas plus two holes is three crossings, which
    # even-odd reads as inside - so the overlap comes back turned over while the two
    # pictures around it do not.
    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, W, H, stroke=0, fill=1)
    c.setFillColorRGB(20 / 255, 20 / 255, 20 / 255)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(30, H - 50, "Two photographs, overlapping")

    under = Image.new("RGB", (180, 130), (206, 44, 30))
    over = Image.new("RGB", (180, 130), (28, 82, 196))
    c.drawImage(ImageReader(under), 30, H - 260, width=180, height=130)
    c.drawImage(ImageReader(over), 130, H - 320, width=180, height=130)
    c.showPage()
    c.save()
    written(OUT / "colours.pdf")

    # encrypted.pdf: the standard security handler, which is what nearly every
    # protected PDF in circulation uses and the only kind pdf.js can unlock.
    enc = StandardEncryption("gander", canPrint=1)
    c = new(OUT / "encrypted.pdf", encrypt=enc)
    c.setFont("Helvetica", 18)
    c.drawString(72, 700, "Unlocked with the password gander.")
    c.showPage()
    c.save()
    written(OUT / "encrypted.pdf")

    # Named .pdf, is not one. pdf.js must say so rather than showing nothing.
    (OUT / "not-a-pdf.pdf").write_bytes(
        b"This file is named .pdf and is plain text. It is not a PDF at all.\n"
    )
    written(OUT / "not-a-pdf.pdf")


# ---------------------------------------------------------------------------
# OOXML
# ---------------------------------------------------------------------------

def docx() -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("Field Survey, Willowmere", level=1)
    doc.add_paragraph(
        "A short report written only so that a test has something to render."
    )
    # A Wingdings bullet sitting in the private use area. docx.html rewrites
    # U+F000 to U+F0FF into real Unicode, because the font is not on the phone
    # and the glyph would otherwise come out as a blank box.
    p = doc.add_paragraph()
    run = p.add_run("")
    run.font.name = "Wingdings"
    run.font.size = Pt(12)
    p.add_run(" A bullet that arrives as a private use codepoint.")
    doc.add_paragraph("The paragraph after it, so ordering is checkable.")
    fix_core_properties(doc)
    doc.save(str(OUT / "report.docx"))
    normalize_zip(OUT / "report.docx")
    written(OUT / "report.docx")


SHEET_ROWS = [
    ("Item", "Quarter", "Amount"),
    ("Surveying", "Q3", 4200),
    ("Drainage", "Q3", 1850),
    ("Fencing", "Q3", 990),
]


def xlsx() -> None:
    import csv

    from openpyxl import Workbook

    wb = Workbook()
    first = wb.active
    first.title = "Summary"
    for row in SHEET_ROWS:
        first.append(row)
    second = wb.create_sheet("Detail")
    second.append(("Note", "Value"))
    second.append(("Second sheet marker", "detail-sheet"))
    third = wb.create_sheet("Notes")
    third.append(("Third sheet marker",))
    wb.properties.created = EPOCH
    wb.properties.modified = EPOCH
    wb.properties.creator = "Gander tests"
    wb.save(str(OUT / "budget.xlsx"))
    normalize_zip(OUT / "budget.xlsx")
    written(OUT / "budget.xlsx")

    with open(OUT / "budget.csv", "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(SHEET_ROWS)
    written(OUT / "budget.csv")


def pptx() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    titles = ["Willowmere Kickoff", "What we found", "What happens next"]
    for n, title in enumerate(titles, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = f"Slide {n} of three."
        body.add_paragraph().text = "Invented content, for rendering only."
    fix_core_properties(prs)
    prs.save(str(OUT / "deck.pptx"))
    normalize_zip(OUT / "deck.pptx")
    written(OUT / "deck.pptx")


# ---------------------------------------------------------------------------
# Text
# ---------------------------------------------------------------------------

def texts() -> None:
    md = """# Willowmere site notes

A heading, a [link](https://example.invalid/notes), and a list:

- first
- second

<script>window.__xss = 1;</script>

<img src="x" onerror="window.__xss = 2;">

Text after the injected markup, so the sanitiser can be seen to have kept it.
"""
    (OUT / "notes.md").write_text(md, encoding="utf-8")
    written(OUT / "notes.md")

    plain = (
        "Plain text, opened by the text viewer.\n"
        "A second line so the newline handling is visible.\n"
        "An accented character: café.\n"
    )
    (OUT / "plain.txt").write_text(plain, encoding="utf-8")
    written(OUT / "plain.txt")

    # Issue #32. A .json file as they nearly always arrive: no spacing at all,
    # the whole document on one line. The id is 2**53 + 1, the smallest whole
    # number a double cannot hold, so a viewer that reformatted by parsing and
    # restringifying would show 9007199254740992 here and be caught. The text
    # carries an accent and an emoji, which is a surrogate pair, because both
    # sit inside a string the formatter has to copy through untouched.
    snapshot = (
        '{"id":9007199254740993,"app":"com.example.reader",'
        '"screen":{"width":1080,"height":2376},"landscape":false,"tags":[],'
        '"nodes":[{"id":0,"parent":-1,"name":"android.widget.FrameLayout",'
        '"text":null,"visible":true},{"id":1,"parent":0,'
        '"name":"android.widget.TextView","text":"café \U0001f600",'
        '"visible":true}],"notes":{}}'
    )
    (OUT / "snapshot.json").write_text(snapshot, encoding="utf-8")
    written(OUT / "snapshot.json")

    # Byte order marks. app.js sniffs these three bytes and picks the decoder;
    # the decoder strips the mark itself, so neither file should show one.
    marked = "Byte order marked text, decoded by the mark alone.\n"
    (OUT / "utf16le.txt").write_bytes(b"\xff\xfe" + marked.encode("utf-16-le"))
    written(OUT / "utf16le.txt")
    (OUT / "utf16be.txt").write_bytes(b"\xfe\xff" + marked.encode("utf-16-be"))
    written(OUT / "utf16be.txt")

    # An extension nothing claims, so the viewer offers to read it as text.
    (OUT / "unknown.xyz").write_bytes(bytes(range(32, 127)) * 4 + b"\n")
    written(OUT / "unknown.xyz")



# ---------------------------------------------------------------------------
# Images pdf.js decodes in WebAssembly
# ---------------------------------------------------------------------------

# Issue #24. Since pdf.js 4 three image encodings are decoded by wasm modules
# the worker fetches on demand rather than by JavaScript in the bundle, and the
# only way to say where those modules are is the wasmUrl option. Gander shipped
# neither the binaries nor the option until 1.17, so every image in these three
# encodings was dropped: the worker warns to a console nobody reads, returns
# nothing, and the page renders with a hole where the picture goes. No error, no
# placeholder, nothing that looks like a failure rather than a layout.
#
# One module, jbig2.wasm, serves both JBIG2 and CCITT fax. CCITT is the ordinary
# compression for a scanned black and white page and by far the commoner of the
# two, and it had been broken the whole time without anyone reporting it, which
# is the argument for testing all three rather than the one that was reported.
#
# Written by hand rather than through ReportLab because the point of each file
# is its /Filter, and a library that re-encodes to something it prefers would
# quietly test nothing.


def _raw_image_pdf(path: Path, w: int, h: int, entries: bytes, data: bytes) -> None:
    """A one-page PDF whose only content is a single image XObject, undecoded."""
    content = b"q 480 0 0 200 20 30 cm /Im0 Do Q"
    objs = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        3: (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 520 260] "
            b"/Resources << /XObject << /Im0 5 0 R >> >> /Contents 4 0 R >>"),
        4: b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        5: (b"<< /Type /XObject /Subtype /Image /Width %d /Height %d "
            b"%s /Length %d >>\nstream\n" % (w, h, entries, len(data))
            + data + b"\nendstream"),
    }
    out = bytearray(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    offsets = {}
    for n in sorted(objs):
        offsets[n] = len(out)
        out += b"%d 0 obj\n" % n + objs[n] + b"\nendobj\n"
    start = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for n in sorted(objs):
        out += b"%010d 00000 n \n" % offsets[n]
    out += (b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
            % (len(objs) + 1, start))
    path.write_bytes(bytes(out))


def _bitonal(w: int, h: int):
    """Black on white, with a circle and two text runs, so a partial decode shows."""
    from PIL import Image, ImageDraw
    im = Image.new("1", (w, h), 1)
    d = ImageDraw.Draw(im)
    d.rectangle([10, 10, w - 11, h - 11], outline=0, width=3)
    d.text((40, 60), "SCANNED PAGE", fill=0)
    d.text((40, 90), "this line proves the decoder ran", fill=0)
    d.ellipse([w - 150, 40, w - 50, 140], outline=0, width=4)
    return im


def _group4(im) -> bytes:
    """T.6 data, taken out of a Group 4 TIFF because Pillow will write one."""
    from PIL import Image
    buf = io.BytesIO()
    im.save(buf, format="TIFF", compression="group4")
    buf.seek(0)
    tiff = Image.open(buf)
    raw = buf.getvalue()
    return b"".join(raw[o:o + c] for o, c in
                    zip(tiff.tag_v2[273], tiff.tag_v2[279]))


def wasm_decoded_images() -> None:
    from PIL import Image, ImageDraw

    W, H = 480, 200

    # jpx.pdf: JPEG 2000, which is what issue #24's file was made of. A
    # photograph rather than a diagram, because JPEG 2000 is a photographic
    # codec and a flat drawing would compress to something unrepresentative.
    photo = Image.new("RGB", (W, H))
    px = photo.load()
    for y in range(H):
        for x in range(W):
            px[x, y] = ((x * 255) // W, (y * 255) // H, ((x + y) * 127) // (W + H))
    d = ImageDraw.Draw(photo)
    d.ellipse([W - 150, 40, W - 50, 140], fill=(250, 250, 40))
    d.text((40, 90), "JPEG 2000", fill=(255, 255, 255))
    buf = io.BytesIO()
    photo.save(buf, format="JPEG2000", irreversible=True, quality_layers=[40])
    # No /ColorSpace: for JPXDecode the codestream carries it, and naming one
    # here would let a reader that ignored the image still look correct.
    _raw_image_pdf(OUT / "jpx.pdf", W, H, b"/Filter /JPXDecode", buf.getvalue())
    written(OUT / "jpx.pdf")

    bitonal = _bitonal(W, H)
    mmr = _group4(bitonal)

    # jbig2.pdf: an embedded JBIG2 stream, which is a bare segment sequence with
    # no file header. A generic region with MMR=1 carries plain T.6 data, so the
    # Group 4 bytes above can be reused and the fixture needs no JBIG2 encoder
    # on the machine generating it.
    page_info = struct.pack(">IIII", W, H, 0, 0) + bytes([0x01]) + struct.pack(">H", 0)
    region = (struct.pack(">IIII", W, H, 0, 0)   # region position and size
              + bytes([0x00])                    # combine into the page by OR
              + bytes([0x01])                    # MMR = 1, so no arithmetic coder
              + mmr)

    def segment(number: int, kind: int, data: bytes) -> bytes:
        return (struct.pack(">I", number)
                + bytes([kind & 0x3F])   # one-byte page association
                + bytes([0x00])          # refers to no other segment
                + bytes([0x01])          # page 1
                + struct.pack(">I", len(data))
                + data)

    jb2 = segment(0, 48, page_info) + segment(1, 39, region)
    _raw_image_pdf(OUT / "jbig2.pdf", W, H,
                   b"/ColorSpace /DeviceGray /BitsPerComponent 1 /Filter /JBIG2Decode",
                   jb2)
    written(OUT / "jbig2.pdf")

    # ccitt.pdf: the same bitmap as plain Group 4, which is what a scanner or a
    # fax produces and what jbig2.wasm turns out to decode as well.
    _raw_image_pdf(OUT / "ccitt.pdf", W, H,
                   b"/ColorSpace /DeviceGray /BitsPerComponent 1 "
                   b"/Filter /CCITTFaxDecode "
                   b"/DecodeParms << /K -1 /Columns %d /Rows %d /BlackIs1 true >>"
                   % (W, H),
                   mmr)
    written(OUT / "ccitt.pdf")


# ---------------------------------------------------------------------------
# Images and audio
# ---------------------------------------------------------------------------

# The six EXIF orientations Thumbs.exifRotation maps to a rotation, plus the
# two flips it folds into 90 and 270.
EXIF_ORIENTATIONS = {
    1: 0,    # normal
    3: 180,  # rotate 180
    6: 90,   # rotate 90
    8: 270,  # rotate 270
    5: 90,   # transpose
    7: 270,  # transverse
}


def images() -> None:
    from PIL import Image, ImageDraw

    def asymmetric(w=120, h=80):
        """A frame with one filled corner, so a rotation is visible."""
        img = Image.new("RGB", (w, h), (245, 245, 245))
        d = ImageDraw.Draw(img)
        d.rectangle([0, 0, w - 1, h - 1], outline=(30, 30, 30), width=2)
        d.rectangle([4, 4, 34, 24], fill=(178, 45, 24))
        return img

    for orientation in sorted(EXIF_ORIENTATIONS):
        path = OUT / f"exif-{orientation}.jpg"
        exif = Image.Exif()
        exif[0x0112] = orientation
        asymmetric().save(path, "JPEG", quality=88, exif=exif)
        written(path)

    asymmetric(64, 64).save(OUT / "tiny.png", "PNG", optimize=True)
    written(OUT / "tiny.png")

    frames = []
    for shift in range(4):
        img = Image.new("P", (48, 48), 0)
        d = ImageDraw.Draw(img)
        d.rectangle([shift * 8, 8, shift * 8 + 16, 24], fill=1)
        img.putpalette([245, 245, 245, 178, 45, 24] + [0] * 762)
        frames.append(img)
    frames[0].save(
        OUT / "anim.gif", save_all=True, append_images=frames[1:],
        duration=120, loop=0,
    )
    written(OUT / "anim.gif")

    (OUT / "icon.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="120" height="80" '
        'viewBox="0 0 120 80" role="img" aria-label="A red square in a frame">\n'
        '  <rect width="120" height="80" fill="#f5f5f5" stroke="#1e1e1e" '
        'stroke-width="2"/>\n'
        '  <rect x="8" y="8" width="30" height="20" fill="#b22d18"/>\n'
        "</svg>\n",
        encoding="utf-8",
    )
    written(OUT / "icon.svg")


def audio() -> None:
    rate, seconds, freq = 8000, 1, 440.0
    frames = bytearray()
    for i in range(rate * seconds):
        # A quiet sine, so a device test that actually plays it is bearable
        frames += struct.pack("<h", int(6000 * sin(2 * pi * freq * i / rate)))
    with wave.open(str(OUT / "tone.wav"), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(rate)
        w.writeframes(bytes(frames))
    written(OUT / "tone.wav")


# ---------------------------------------------------------------------------
# Zips, issue #30
# ---------------------------------------------------------------------------

# Written by hand, like the wasm PDFs above and for the same reason: the point of
# most of these is a byte Python's zipfile will not write. It sets the UTF-8 flag
# on every name that is not ASCII, and the names Gander has to decode are exactly
# the ones a Windows machine writes in its own code page with that flag clear.

ZIP_DOS_DATE = ((ZIP_DATE[0] - 1980) << 9) | (ZIP_DATE[1] << 5) | ZIP_DATE[2]
HOST_DOS, HOST_UNIX = 0, 3
STORED, DEFLATED, DEFLATE64 = 0, 8, 9
AES = 99


class Member:
    """One entry. A method other than stored, deflated or Deflate64 writes data as given.

    zipcrypto or aes, a password, encrypts it: aes as (strength, version), strength 1 to 3
    for 128 to 256 bit keys and version 1 or 2 for AE-1 or AE-2.
    """

    def __init__(self, name: bytes, data: bytes = b"", *, method=DEFLATED, flags=0,
                 host=HOST_UNIX, extra=b"", descriptor=False, directory=False, time=0,
                 zipcrypto: bytes = None, aes=None, password: bytes = None):
        self.name, self.data, self.method = name, data, method
        self.flags = flags | (0x08 if descriptor else 0) | (0x01 if zipcrypto or aes else 0)
        self.host, self.extra = host, extra
        self.descriptor, self.directory, self.time = descriptor, directory, time
        self.zipcrypto, self.aes, self.password = zipcrypto, aes, password


def _extra(field_id: int, data: bytes) -> bytes:
    return struct.pack("<HH", field_id, len(data)) + data


# PKWARE's original encryption, APPNOTE 6.1. Twelve bytes in front, the last of them the
# top byte of the CRC, or of the time when the CRC trails the data; the eleven before it
# would be random, and are a hash of the name here so the output never changes.

def _crc_byte(crc: int, b: int) -> int:
    return zlib.crc32(bytes([b]), crc ^ 0xFFFFFFFF) ^ 0xFFFFFFFF


class _ZipCrypto:
    def __init__(self, password: bytes):
        self.k = [0x12345678, 0x23456789, 0x34567890]
        for b in password:
            self._update(b)

    def _update(self, b: int) -> None:
        k0 = _crc_byte(self.k[0], b)
        k1 = ((self.k[1] + (k0 & 0xFF)) * 134775813 + 1) & 0xFFFFFFFF
        self.k = [k0, k1, _crc_byte(self.k[2], k1 >> 24)]

    def encrypt(self, data: bytes) -> bytes:
        out = bytearray()
        for b in data:
            t = (self.k[2] | 2) & 0xFFFF
            out.append(b ^ (((t * (t ^ 1)) >> 8) & 0xFF))
            self._update(b)
        return bytes(out)


def zipcrypto(body: bytes, password: bytes, check: int, name: bytes) -> bytes:
    keys = _ZipCrypto(password)
    header = hashlib.sha256(b"header" + name).digest()[:11] + bytes([check])
    return keys.encrypt(header) + keys.encrypt(body)


# WinZip's AES, from its published specification: PBKDF2-HMAC-SHA1 over the password and a
# salt, a thousand rounds, into the AES key, an HMAC key and a two byte verifier; AES in
# counter mode counting up from one little-endian; the first ten bytes of HMAC-SHA1 over what
# was encrypted, after it. The salt is a hash of the name, again so nothing changes.

def winzip_aes(body: bytes, password: bytes, strength: int, name: bytes) -> bytes:
    from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
    size = {1: 16, 2: 24, 3: 32}[strength]
    salt = hashlib.sha256(b"salt" + name).digest()[:size // 2]
    keys = hashlib.pbkdf2_hmac("sha1", password, salt, 1000, 2 * size + 2)
    aes = Cipher(algorithms.AES(keys[:size]), modes.ECB()).encryptor()
    out = bytearray()
    for n, at in enumerate(range(0, len(body), 16), start=1):
        stream = aes.update(n.to_bytes(16, "little"))
        out += bytes(a ^ b for a, b in zip(body[at:at + 16], stream))
    code = hmac.new(keys[size:2 * size], bytes(out), hashlib.sha1).digest()[:10]
    return salt + keys[2 * size:] + bytes(out) + code


# Deflate64, which zlib cannot write. Enough of an encoder to put every part of the format
# in one stream: a stored block, a fixed one and dynamic ones, matches reaching into the
# second 32 KB of the window, and lengths past 258, which only Deflate64's last length code
# can carry. Greedy matching over the whole 64 KB window, nothing cleverer.

_LEN_BASE = [3, 4, 5, 6, 7, 8, 9, 10, 11, 13, 15, 17, 19, 23, 27, 31, 35, 43, 51, 59, 67, 83,
             99, 115, 131, 163, 195, 227, 3]
_LEN_EXTRA = [0, 0, 0, 0, 0, 0, 0, 0, 1, 1, 1, 1, 2, 2, 2, 2, 3, 3, 3, 3, 4, 4, 4, 4, 5, 5, 5,
              5, 16]
_DIST_BASE = [1, 2, 3, 4, 5, 7, 9, 13, 17, 25, 33, 49, 65, 97, 129, 193, 257, 385, 513, 769,
              1025, 1537, 2049, 3073, 4097, 6145, 8193, 12289, 16385, 24577, 32769, 49153]
_DIST_EXTRA = [0, 0, 0, 0, 1, 1, 2, 2, 3, 3, 4, 4, 5, 5, 6, 6, 7, 7, 8, 8, 9, 9, 10, 10, 11,
               11, 12, 12, 13, 13, 14, 14]
_ORDER = [16, 17, 18, 0, 8, 7, 9, 6, 10, 5, 11, 4, 12, 3, 13, 2, 14, 1, 15]


class _Bits:
    def __init__(self):
        self.out, self.acc, self.n = bytearray(), 0, 0

    def put(self, value: int, count: int) -> None:
        self.acc |= value << self.n
        self.n += count
        while self.n >= 8:
            self.out.append(self.acc & 0xFF)
            self.acc >>= 8
            self.n -= 8

    def code(self, code: int, length: int) -> None:
        self.put(int(format(code, f"0{length}b")[::-1], 2), length)

    def align(self) -> None:
        if self.n:
            self.out.append(self.acc & 0xFF)
        self.acc = self.n = 0


def _lengths(freqs, limit):
    used = [(f, s) for s, f in enumerate(freqs) if f]
    lengths = [0] * len(freqs)
    if not used:
        # A block of nothing but literals has no distance code at all, which is allowed
        return lengths
    if len(used) == 1:
        # One code alone is incomplete; a second, never used, completes it
        lengths[used[0][1]] = 1
        lengths[0 if used[0][1] else 1] = 1
        return lengths
    while True:
        heap = [(f, i, (s,)) for i, (f, s) in enumerate(used)]
        heapq.heapify(heap)
        depth, tie = {s: 0 for _, s in used}, len(heap)
        while len(heap) > 1:
            f1, _, a = heapq.heappop(heap)
            f2, _, b = heapq.heappop(heap)
            for s in a + b:
                depth[s] += 1
            heapq.heappush(heap, (f1 + f2, tie, a + b))
            tie += 1
        if max(depth.values()) <= limit:
            for s, d in depth.items():
                lengths[s] = d
            return lengths
        used = [((f + 1) // 2, s) for f, s in used]


def _canonical(lengths):
    count = [0] * 16
    for n in lengths:
        if n:
            count[n] += 1
    code, first = 0, [0] * 16
    for n in range(1, 16):
        code = (code + count[n - 1]) << 1
        first[n] = code
    codes = []
    for n in lengths:
        codes.append(first[n])
        if n:
            first[n] += 1
    return codes


def _length_code(n):
    if n > 258:
        return 285, n - 3, 16
    i = max(i for i in range(28) if _LEN_BASE[i] <= n)
    return 257 + i, n - _LEN_BASE[i], _LEN_EXTRA[i]


def _distance_code(d):
    i = max(i for i in range(32) if _DIST_BASE[i] <= d)
    return i, d - _DIST_BASE[i], _DIST_EXTRA[i]


def _matches(data: bytes, start: int):
    """Greedy LZ77 over a 64 KB window, from start, reaching back before it."""
    table, tokens, i = {}, [], 0
    for j in range(max(0, start - 65536), start):
        table.setdefault(data[j:j + 3], []).append(j)
    i = start
    while i < len(data):
        best, where = 0, 0
        for j in reversed(table.get(data[i:i + 3], [])[-32:]):
            if i - j > 65536:
                break
            n = 0
            while i + n < len(data) and n < 65538 and data[j + n] == data[i + n]:
                n += 1
            if n > best:
                best, where = n, i - j
        step = best if best >= 3 else 1
        for j in range(i, min(i + step, len(data) - 2)):
            table.setdefault(data[j:j + 3], []).append(j)
        tokens.append((best, where) if best >= 3 else data[i])
        i += step
    return tokens


def _block(bits: _Bits, tokens, last: bool, dynamic: bool) -> None:
    if dynamic:
        lit, dist = [0] * 286, [0] * 32
        lit[256] = 1
        for t in tokens:
            if isinstance(t, int):
                lit[t] += 1
            else:
                lit[_length_code(t[0])[0]] += 1
                dist[_distance_code(t[1])[0]] += 1
        lit_lengths, dist_lengths = _lengths(lit, 15), _lengths(dist, 15)
        hlit = max(257, max(s for s, n in enumerate(lit_lengths) if n) + 1)
        hdist = max(1, max((s for s, n in enumerate(dist_lengths) if n), default=0) + 1)
        seq, all_lengths, i = [], lit_lengths[:hlit] + dist_lengths[:hdist], 0
        while i < len(all_lengths):
            n = all_lengths[i]
            run = 1
            while i + run < len(all_lengths) and all_lengths[i + run] == n:
                run += 1
            if n == 0 and run >= 3:
                r = min(run, 138)
                seq.append((18, r - 11, 7) if r >= 11 else (17, r - 3, 3))
                i += r
            elif n and run >= 4:
                r = min(run - 1, 6)
                seq += [(n, 0, 0), (16, r - 3, 2)]
                i += 1 + r
            else:
                seq.append((n, 0, 0))
                i += 1
        cl = [0] * 19
        for sym, _, _ in seq:
            cl[sym] += 1
        cl_lengths = _lengths(cl, 7)
        hclen = 19
        while hclen > 4 and cl_lengths[_ORDER[hclen - 1]] == 0:
            hclen -= 1
        bits.put(1 if last else 0, 1)
        bits.put(2, 2)
        bits.put(hlit - 257, 5)
        bits.put(hdist - 1, 5)
        bits.put(hclen - 4, 4)
        for i in range(hclen):
            bits.put(cl_lengths[_ORDER[i]], 3)
        cl_codes = _canonical(cl_lengths)
        for sym, value, extra in seq:
            bits.code(cl_codes[sym], cl_lengths[sym])
            if extra:
                bits.put(value, extra)
    else:
        lit_lengths = [8] * 144 + [9] * 112 + [7] * 24 + [8] * 8
        dist_lengths = [5] * 32
        bits.put(1 if last else 0, 1)
        bits.put(1, 2)
    lit_codes, dist_codes = _canonical(lit_lengths), _canonical(dist_lengths)
    for t in tokens:
        if isinstance(t, int):
            bits.code(lit_codes[t], lit_lengths[t])
        else:
            sym, value, extra = _length_code(t[0])
            bits.code(lit_codes[sym], lit_lengths[sym])
            bits.put(value, extra)
            sym, value, extra = _distance_code(t[1])
            bits.code(dist_codes[sym], dist_lengths[sym])
            bits.put(value, extra)
    bits.code(lit_codes[256], lit_lengths[256])


def deflate64(data: bytes) -> bytes:
    """A stored block, then a fixed one, then two dynamic ones."""
    bits = _Bits()
    stored = data[:min(len(data), 16 * 1024)]
    bits.put(0, 1)
    bits.put(0, 2)
    bits.align()
    bits.out += struct.pack("<HH", len(stored), len(stored) ^ 0xFFFF) + stored
    tokens = _matches(data, len(stored))
    fixed, rest = tokens[:500], tokens[500:]
    half = len(rest) // 2
    _block(bits, fixed, last=False, dynamic=False)
    _block(bits, rest[:half], last=False, dynamic=True)
    _block(bits, rest[half:], last=True, dynamic=True)
    bits.align()
    return bytes(bits.out)


def zip_bytes(members, *, zip64=False, comment=b"") -> bytes:
    out, central = bytearray(), bytearray()
    for m in members:
        crc = zlib.crc32(m.data) & 0xFFFFFFFF
        if m.method == DEFLATED:
            packer = zlib.compressobj(9, zlib.DEFLATED, -15)
            body = packer.compress(m.data) + packer.flush()
        elif m.method == DEFLATE64:
            body = deflate64(m.data)
        else:
            body = m.data
        method, aes_extra = m.method, b""
        if m.zipcrypto:
            check = (m.time >> 8) if m.descriptor else (crc >> 24)
            body = zipcrypto(body, m.zipcrypto, check, m.name)
        elif m.aes:
            strength, version = m.aes
            body = winzip_aes(body, m.password, strength, m.name)
            aes_extra = _extra(0x9901, struct.pack("<H2sBH", version, b"AE", strength, m.method))
            method = AES
            if version == 2:
                crc = 0
        if method == AES:
            needed = 51
        elif method == DEFLATE64:
            needed = 21
        else:
            needed = 63 if method not in (STORED, DEFLATED) else (45 if zip64 else 20)
        made_by = (m.host << 8) | (45 if zip64 else 20)
        if m.host == HOST_UNIX:
            external = ((0o40755 << 16) | 0x10) if m.directory else (0o100644 << 16)
        else:
            external = 0x10 if m.directory else 0x20
        offset = len(out)

        local_crc, local_c, local_u = (0, 0, 0) if m.descriptor else (crc, len(body), len(m.data))
        local_extra = aes_extra
        if zip64:
            local_extra = _extra(1, struct.pack("<QQ", len(m.data), len(body))) + local_extra
            local_c = local_u = 0xFFFFFFFF
        out += struct.pack("<IHHHHHIIIHH", 0x04034B50, needed, m.flags, method, m.time,
                           ZIP_DOS_DATE, local_crc, local_c, local_u, len(m.name),
                           len(local_extra))
        out += m.name + local_extra + body
        if m.descriptor:
            out += struct.pack("<IIII", 0x08074B50, crc, len(body), len(m.data))

        central_extra = m.extra + aes_extra
        c_size, u_size, c_offset = len(body), len(m.data), offset
        if zip64:
            central_extra = _extra(1, struct.pack("<QQQ", len(m.data), len(body), offset)) + central_extra
            c_size = u_size = c_offset = 0xFFFFFFFF
        central += struct.pack("<IHHHHHHIIIHHHHHII", 0x02014B50, made_by, needed, m.flags,
                               method, m.time, ZIP_DOS_DATE, crc, c_size, u_size, len(m.name),
                               len(central_extra), 0, 0, 0, external, c_offset)
        central += m.name + central_extra

    index_at = len(out)
    out += central
    count = len(members)
    if zip64:
        record_at = len(out)
        out += struct.pack("<IQHHIIQQQQ", 0x06064B50, 44, 45, 45, 0, 0, count, count,
                           len(central), index_at)
        out += struct.pack("<IIQI", 0x07064B50, 0, record_at, 1)
        out += struct.pack("<IHHHHIIH", 0x06054B50, 0, 0, 0xFFFF, 0xFFFF,
                           0xFFFFFFFF, 0xFFFFFFFF, len(comment))
    else:
        out += struct.pack("<IHHHHIIH", 0x06054B50, 0, 0, count, count, len(central),
                           index_at, len(comment))
    return bytes(out + comment)


def zips() -> None:
    pdf = (OUT / "six-pages.pdf").read_bytes()
    png = (OUT / "tiny.png").read_bytes()
    notes = (OUT / "notes.md").read_bytes()
    plain = (OUT / "plain.txt").read_bytes()
    inner = zip_bytes([Member(b"inside.txt", b"A file inside a zip inside a zip.\n")])
    # Extended timestamp, 2026-01-01T00:00:00Z: the one entry whose time has a zone
    utc = _extra(0x5455, struct.pack("<BI", 1, 1767225600))

    # What people actually have: folders with and without their own entries, a
    # compressed PDF, a stored photo, a file whose sizes trail its data, the
    # clutter macOS adds, zips inside the zip both ways, and the two kinds of file
    # that are listed and cannot be opened. Plus a comment, which moves the end
    # record away from the end of the file.
    archive = [
        Member(b"reports/", method=STORED, directory=True),
        Member(b"reports/six-pages.pdf", pdf, extra=utc),
        Member(b"reports/notes.md", notes, descriptor=True),
        Member(b"photos/tiny.png", png, method=STORED),
        Member(b"photos/.hidden-thumbs", b"not for people"),
        Member(b"__MACOSX/photos/._tiny.png", b"resource fork"),
        Member(b"plain.txt", plain),
        Member(b"nested/stored.zip", inner, method=STORED),
        Member(b"nested/packed.zip", inner),
        Member(b"private/locked.txt", b"The password was gander.\n", zipcrypto=b"gander"),
        Member(b"private/table.dat", bytes(range(40)), method=14),
    ]
    (OUT / "archive.zip").write_bytes(zip_bytes(archive, comment=b"Gander test archive"))
    written(OUT / "archive.zip")

    # Names that would climb out of a folder, start at a root, or use Windows'
    # separator, and one that reorders itself to look like another file.
    odd = [
        Member(b"../escaped.txt", b"one"),
        Member(b"/absolute/path.txt", b"two"),
        Member(b"docs\\readme.txt", b"three", host=HOST_DOS),
        Member(b"unix\\name.txt", b"four"),
        Member(b"a//b/./c.txt", b"five"),
        Member(b"dup.txt", b"first"),
        Member(b"dup.txt", b"second"),
        Member("photo\u202Egpj.apk".encode(), b"six", flags=0x800),
    ]
    (OUT / "odd-names.zip").write_bytes(zip_bytes(odd))
    written(OUT / "odd-names.zip")

    (OUT / "zip64.zip").write_bytes(zip_bytes(
        [Member(b"big/report.pdf", pdf), Member(b"big/notes.txt", plain)], zip64=True))
    written(OUT / "zip64.zip")

    # Every way a file can be under a password that Gander reads, and the one it does not,
    # all with the password gander, beside a file with none. A half past two in the
    # afternoon on the file whose CRC trails it, since that is what its check byte is.
    half_past_two = (14 << 11) | (30 << 5)
    (OUT / "locked.zip").write_bytes(zip_bytes([
        Member(b"open.txt", plain),
        Member(b"zipcrypto.txt", plain, zipcrypto=b"gander"),
        Member(b"zipcrypto-trailing.md", notes, zipcrypto=b"gander", descriptor=True,
               time=half_past_two),
        Member(b"zipcrypto.png", png, method=STORED, zipcrypto=b"gander"),
        Member(b"aes128.txt", plain, aes=(1, 1), password=b"gander"),
        Member(b"aes192.pdf", pdf, aes=(2, 2), password=b"gander"),
        Member(b"aes256.png", png, method=STORED, aes=(3, 2), password=b"gander"),
        Member(b"aes256-deflate64.md", notes, method=DEFLATE64, aes=(3, 2), password=b"gander"),
        Member(b"strong.bin", bytes(range(64)), method=STORED, flags=0x41),
    ]))
    written(OUT / "locked.zip")

    # A password that is not ASCII, as each encryption takes it: WinZip's AES as UTF-8, and
    # the older one as a Russian Windows machine's own code page.
    (OUT / "locked-cyrillic.zip").write_bytes(zip_bytes([
        Member(b"zipcrypto.txt", plain, zipcrypto="пароль".encode("cp866")),
        Member(b"aes.txt", plain, aes=(3, 2), password="пароль".encode("utf-8")),
    ]))
    written(OUT / "locked-cyrillic.zip")

    # Deflate64, which Windows writes for anything over 2 GB. Text to compress, then the
    # three things only Deflate64 has: a match from more than 32 KB back, one from more than
    # 48 KB back, and matches far longer than 258.
    rng = random.Random(64)
    words = [w.encode() for w in (plain + notes).decode().split() if w.isalpha()]
    prose = b" ".join(rng.choice(words) for _ in range(12000))[:70000]
    long_text = (prose + prose[30000:33000] + prose[5000:6000] + b"ab" * 3000 +
                 prose[10000:30000])
    (OUT / "deflate64.zip").write_bytes(zip_bytes([
        Member(b"long.txt", long_text, method=DEFLATE64),
        Member(b"short.txt", plain, method=DEFLATE64),
        Member(b"empty.txt", b"", method=DEFLATE64),
    ]))
    written(OUT / "deflate64.zip")

    # Names as Windows writes them in each code page, flag clear, and as macOS
    # writes them, UTF-8 with the flag clear too.
    text = b"Plain text inside a zip.\n"
    for fixture, encoding, names in [
        ("names-gbk.zip", "gbk", ["季度报告/会议记录.txt", "照片/北京旅行.txt"]),
        ("names-cp866.zip", "cp866", ["Документы/Отчёт за квартал.txt", "Фото/Москва.txt"]),
        ("names-sjis.zip", "shift_jis", ["資料/報告書.txt", "資料/議事録.txt"]),
        ("names-korean.zip", "cp949", ["문서/분기 보고서.pdf", "사진/제주도 여행.jpg"]),
        ("names-mac.zip", "utf-8", ["Отчёт/报告 résumé.txt"]),
    ]:
        members = [Member(n.encode(encoding), text) for n in names]
        (OUT / fixture).write_bytes(zip_bytes(members))
        written(OUT / fixture)


# ---------------------------------------------------------------------------


# ---------------------------------------------------------------------------
# The three word processor formats prose.html reads itself: OpenDocument text,
# Rich Text and Word 97-2003. Each is written here by hand, from the format,
# rather than by a library, so that the bytes are stable and say exactly what
# the tests rely on. What each one carries is the same short document.
#
# These prove the page's plumbing: order, formatting, tables, pictures, notes,
# the error card. They cannot prove the readers against the world, because a
# file written from the same understanding of a format as the reader agrees
# with it by construction; that was done against LibreOffice's rendering of
# real files from the Apache POI and Tika corpora when the readers were built.
# ---------------------------------------------------------------------------

PROSE_TITLE = "Field Survey, Willowmere"
PROSE_BODY = "A short report written only so that a test has something to render."
PROSE_BOLD = "Bold words"
PROSE_AFTER = "The paragraph after it, so ordering is checkable."
PROSE_HINDI = "\u092a\u0941\u0932 \u092c\u0902\u0926 \u0939\u0948\u0964"    # the bridge is closed
PROSE_NOTE = "A footnote, which lands under a rule at the end."


def odt() -> None:
    """letter.odt: a zip of XML, the way LibreOffice writes one, with a picture."""
    png = (OUT / "tiny.png").read_bytes()
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">\n'
        ' <manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>\n'
        ' <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>\n'
        ' <manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>\n'
        ' <manifest:file-entry manifest:full-path="Pictures/tiny.png" manifest:media-type="image/png"/>\n'
        '</manifest:manifest>\n'
    )
    ns = (
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
        'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
        'xmlns:xlink="http://www.w3.org/1999/xlink" '
        'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
        'office:version="1.3"'
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<office:document-styles {ns}>\n'
        ' <office:font-face-decls>\n'
        '  <style:font-face style:name="Liberation Serif" svg:font-family="\'Liberation Serif\'" style:font-family-generic="roman"/>\n'
        ' </office:font-face-decls>\n'
        ' <office:styles>\n'
        '  <style:default-style style:family="paragraph">\n'
        '   <style:text-properties style:font-name="Liberation Serif" fo:font-size="12pt"/>\n'
        '  </style:default-style>\n'
        '  <style:style style:name="Standard" style:family="paragraph"/>\n'
        '  <style:style style:name="Heading_20_1" style:display-name="Heading 1" style:family="paragraph" style:parent-style-name="Standard">\n'
        '   <style:paragraph-properties fo:margin-top="12pt" fo:margin-bottom="6pt"/>\n'
        '   <style:text-properties fo:font-size="18pt" fo:font-weight="bold" fo:color="#1f3864"/>\n'
        '  </style:style>\n'
        ' </office:styles>\n'
        ' <office:automatic-styles>\n'
        '  <style:page-layout style:name="pm1">\n'
        '   <style:page-layout-properties fo:page-width="21cm" fo:page-height="29.7cm" fo:margin-top="2cm" fo:margin-bottom="2cm" fo:margin-left="2cm" fo:margin-right="2cm"/>\n'
        '  </style:page-layout>\n'
        ' </office:automatic-styles>\n'
        ' <office:master-styles>\n'
        '  <style:master-page style:name="Standard" style:page-layout-name="pm1">\n'
        '   <style:header><text:p>HEADER-MARK Willowmere Parish Council</text:p></style:header>\n'
        '  </style:master-page>\n'
        ' </office:master-styles>\n'
        '</office:document-styles>\n'
    )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<office:document-content {ns}>\n'
        ' <office:automatic-styles>\n'
        '  <style:style style:name="T1" style:family="text"><style:text-properties fo:font-weight="bold"/></style:style>\n'
        '  <style:style style:name="P1" style:family="paragraph" style:parent-style-name="Standard">\n'
        '   <style:paragraph-properties fo:break-before="page"/>\n'
        '  </style:style>\n'
        '  <style:style style:name="Table1.A1" style:family="table-cell">\n'
        '   <style:table-cell-properties fo:border="0.5pt solid #000000" fo:background-color="#d9e2f3"/>\n'
        '  </style:style>\n'
        '  <text:list-style style:name="L1">\n'
        '   <text:list-level-style-bullet text:level="1" text:bullet-char="\u2022">\n'
        '    <style:list-level-properties><style:list-level-label-alignment text:label-followed-by="listtab" fo:text-indent="-0.635cm" fo:margin-left="1.27cm"/></style:list-level-properties>\n'
        '   </text:list-level-style-bullet>\n'
        '  </text:list-style>\n'
        ' </office:automatic-styles>\n'
        ' <office:body>\n'
        '  <office:text>\n'
        f'   <text:h text:style-name="Heading_20_1" text:outline-level="1">{PROSE_TITLE}</text:h>\n'
        f'   <text:p text:style-name="Standard">{PROSE_BODY}</text:p>\n'
        f'   <text:p text:style-name="Standard"><text:span text:style-name="T1">{PROSE_BOLD}</text:span> in the middle of a sentence.'
        f'<text:note text:id="ftn1" text:note-class="footnote"><text:note-citation>1</text:note-citation>'
        f'<text:note-body><text:p>{PROSE_NOTE}</text:p></text:note-body></text:note></text:p>\n'
        '   <text:list text:style-name="L1">\n'
        '    <text:list-item><text:p>Bridges and culverts</text:p></text:list-item>\n'
        '    <text:list-item><text:p>Retaining walls</text:p></text:list-item>\n'
        '   </text:list>\n'
        '   <table:table table:name="Table1">\n'
        '    <table:table-column table:number-columns-repeated="2"/>\n'
        '    <table:table-row><table:table-cell table:style-name="Table1.A1"><text:p>Item</text:p></table:table-cell>'
        '<table:table-cell table:style-name="Table1.A1"><text:p>Amount</text:p></table:table-cell></table:table-row>\n'
        '    <table:table-row><table:table-cell><text:p>Surveying</text:p></table:table-cell>'
        '<table:table-cell><text:p>4200</text:p></table:table-cell></table:table-row>\n'
        '   </table:table>\n'
        '   <text:p text:style-name="Standard"><draw:frame draw:name="Picture" text:anchor-type="as-char" svg:width="2cm" svg:height="2cm">'
        '<draw:image xlink:href="Pictures/tiny.png" xlink:type="simple"/></draw:frame></text:p>\n'
        f'   <text:p text:style-name="P1">{PROSE_AFTER}</text:p>\n'
        f'   <text:p text:style-name="Standard">Hindi: {PROSE_HINDI}</text:p>\n'
        '  </office:text>\n'
        ' </office:body>\n'
        '</office:document-content>\n'
    )
    members = [
        Member(b"mimetype", b"application/vnd.oasis.opendocument.text", method=STORED),
        Member(b"META-INF/manifest.xml", manifest.encode()),
        Member(b"styles.xml", styles.encode()),
        Member(b"content.xml", content.encode()),
        Member(b"Pictures/tiny.png", png, method=STORED),
    ]
    (OUT / "letter.odt").write_bytes(zip_bytes(members))
    written(OUT / "letter.odt")


def rtf() -> None:
    """memo.rtf: what Word writes, with a code page, Unicode escapes and a picture."""
    png = (OUT / "tiny.png").read_bytes()

    def uni(text: str) -> str:
        # Each character outside ASCII as \uN, with the one-byte fallback \uc1 asks for
        return "".join(c if ord(c) < 128 else f"\\u{ord(c) if ord(c) < 0x8000 else ord(c) - 0x10000}?" for c in text)

    doc = (
        "{\\rtf1\\ansi\\ansicpg1252\\deff0\\deflang1033\\uc1"
        "{\\fonttbl{\\f0\\froman\\fcharset0 Times New Roman;}{\\f1\\fswiss\\fcharset0 Arial;}{\\f2\\fnil\\fcharset2 Wingdings;}}\n"
        "{\\colortbl;\\red31\\green56\\blue100;\\red255\\green255\\blue0;}\n"
        "{\\stylesheet{\\s0 Normal;}{\\s1\\f1\\fs36\\b\\cf1 heading 1;}}\n"
        "{\\*\\generator Gander fixtures}\n"
        "{\\info{\\title Field Survey}}\n"
        "\\paperw11906\\paperh16838\\margl1134\\margr1134\\margt1134\\margb1134\n"
        "{\\header\\pard\\plain\\f0\\fs20 HEADER-MARK Willowmere Parish Council\\par}\n"
        f"\\pard\\plain\\s1\\f1\\fs36\\b\\cf1 {PROSE_TITLE}\\par\n"
        f"\\pard\\plain\\f0\\fs24 {PROSE_BODY}\\par\n"
        f"\\pard\\plain\\f0\\fs24 {{\\b {PROSE_BOLD}}} in the middle of a sentence."
        f"{{\\super\\chftn{{\\footnote\\pard\\plain\\f0\\fs20 {{\\super\\chftn}} {PROSE_NOTE}\\par}}}}\\par\n"
        "{\\pntext\\f2\\'b7\\tab}{\\*\\pn\\pnlvlblt\\pnf2\\pnindent360{\\pntxtb\\'b7}}\\pard\\plain\\fi-360\\li720 Bridges and culverts\\par\n"
        "{\\pntext\\f2\\'b7\\tab}\\pard\\plain\\fi-360\\li720 Retaining walls\\par\n"
        "\\trowd\\trgaph108\\cellx4320\\cellx8640\\clcbpat2"
        "\\pard\\intbl Item\\cell Amount\\cell\\row\n"
        "\\trowd\\trgaph108\\cellx4320\\cellx8640"
        "\\pard\\intbl Surveying\\cell 4200\\cell\\row\n"
        "\\pard\\plain\\f0\\fs24 {\\pict\\pngblip\\picw64\\pich64\\picwgoal1134\\pichgoal1134 " + png.hex() + "}\\par\n"
        "\\page\n"
        f"\\pard\\plain\\f0\\fs24 {PROSE_AFTER}\\par\n"
        f"\\pard\\plain\\f0\\fs24 Hindi: {uni(PROSE_HINDI)}\\par\n"
        "\\pard\\plain\\f0\\fs24 Accents: {\\'e9}t{\\'e9} and Gr{\\'f6}{\\'df}e.\\par\n"
        "}\n"
    )
    (OUT / "memo.rtf").write_bytes(doc.encode("ascii"))
    written(OUT / "memo.rtf")


def doc() -> None:
    """legacy.doc: a Word 97 file, with its text in UTF-16, a bold run and a table.

    The compound file has one FAT sector and one directory sector, and each of
    its two streams is padded to 4,096 bytes so that neither is small enough to
    need the mini stream.
    """
    cell = "\x07"
    parts = [
        PROSE_TITLE + "\r",
        PROSE_BODY + "\r",
        PROSE_BOLD, " in the middle of a sentence.\r",
        "Item" + cell, "Amount" + cell, cell,
        "Surveying" + cell, "4200" + cell, cell,
        PROSE_AFTER + "\r",
        "Hindi: " + PROSE_HINDI + "\r",
    ]
    text = "".join(parts)
    fc_text = 0x400
    fc_end = fc_text + 2 * len(text)
    assert fc_end <= 0x800

    # The paragraphs, and what each one's properties say: nothing, in a table
    # cell, or the end of a row with the row's own definition. The bold words are
    # a run inside the third paragraph, not a paragraph of their own
    para_runs, fc = [], fc_text
    for part in parts:
        end = fc + 2 * len(part)
        if part != PROSE_BOLD:
            start = fc - (2 * len(PROSE_BOLD) if part.startswith(" in the middle") else 0)
            kind = "row" if part == cell else "cell" if part.endswith(cell) else "plain"
            para_runs.append((start, end, kind))
        fc = end

    bold_start = fc_text + 2 * len(parts[0] + parts[1])
    bold_end = bold_start + 2 * len(PROSE_BOLD)

    def u16(v): return struct.pack("<H", v)
    def u32(v): return struct.pack("<I", v)
    def i16(v): return struct.pack("<h", v)

    # CHPX page: three runs, the middle one bold
    chpx = bytes([3, 0x35, 0x08, 1])                       # cb, sprmCFBold, on
    page = bytearray(512)
    for i, f in enumerate([fc_text, bold_start, bold_end, fc_end]):
        page[i * 4:i * 4 + 4] = u32(f)
    page[500:504] = chpx
    page[16:19] = bytes([0, 250, 0])
    page[511] = 3
    chp_page = bytes(page)

    # PAPX page: one run per paragraph, properties at the end of the page
    in_table = bytes([3, 0, 0, 0x16, 0x24, 1])              # cb 3: istd, sprmPFInTable
    tc = u16(0) + u16(4320) + bytes([4, 1, 0, 0]) * 4      # a cell: half-point single borders
    def_table = bytes([2]) + i16(0) + i16(4320) + i16(8640) + tc + tc
    row_end = (u16(0) + bytes([0x16, 0x24, 1, 0x17, 0x24, 1]) +
               bytes([0x08, 0xD6]) + u16(len(def_table) + 1) + def_table)
    assert len(row_end) % 2 == 1
    row_end = bytes([(len(row_end) + 1) // 2]) + row_end
    page = bytearray(512)
    n = len(para_runs)
    for i, (start, _, _) in enumerate(para_runs):
        page[i * 4:i * 4 + 4] = u32(start)
    page[n * 4:n * 4 + 4] = u32(para_runs[-1][1])
    page[400:400 + len(row_end)] = row_end
    page[460:460 + len(in_table)] = in_table
    for i, (_, _, kind) in enumerate(para_runs):
        at = (n + 1) * 4 + i * 13
        page[at] = {"plain": 0, "cell": 230, "row": 200}[kind]
    page[511] = n
    pap_page = bytes(page)

    # The table stream: the piece table, an empty stylesheet, one font, the bin tables
    clx = bytes([2]) + u32(16) + u32(0) + u32(len(text)) + u16(0) + u32(fc_text) + u16(0)
    stsh = u16(18) + u16(0) + u16(10) + u16(0) * 4 + u16(0) * 3
    name = "Times New Roman".encode("utf-16-le") + b"\0\0"
    ffn = bytes([0, 0x12]) + u16(400) + bytes([0, 0]) + bytes(10) + bytes(24) + name
    ffn = bytes([len(ffn) - 1]) + ffn[1:]
    fonts = u16(1) + u16(0) + ffn
    bte_chpx = u32(fc_text) + u32(0x800) + u32(4)
    bte_papx = u32(fc_text) + u32(0x800) + u32(5)
    table = bytearray(4096)
    places = {}
    at = 0
    for key, blob in (("clx", clx), ("stsh", stsh), ("ffn", fonts), ("chpx", bte_chpx), ("papx", bte_papx)):
        table[at:at + len(blob)] = blob
        places[key] = (at, len(blob))
        at += len(blob) + (16 - len(blob) % 16)

    # The FIB, then the text at fc 0x400, then the two formatting pages
    fib = bytearray(0x400)
    fib[0:2] = u16(0xA5EC)
    fib[2:4] = u16(0xC1)
    fib[6:8] = u16(0x0409)
    fib[0x0A:0x0C] = u16(0x0200)                             # fWhichTblStm: 1Table
    fib[0x0C:0x0E] = u16(0xBF)
    fib[0x18:0x1C] = u32(fc_text)
    fib[0x1C:0x20] = u32(fc_end)
    fib[0x20:0x22] = u16(0x0E)
    fib[0x3E:0x40] = u16(0x16)
    fib[0x40:0x44] = u32(fc_end)
    fib[0x4C:0x50] = u32(len(text))
    fib[0x98:0x9A] = u16(0x5D)
    for index, key in ((1, "stsh"), (12, "chpx"), (13, "papx"), (15, "ffn"), (33, "clx")):
        fib[0x9A + index * 8:0x9A + index * 8 + 8] = u32(places[key][0]) + u32(places[key][1])
    word = bytearray(4096)
    word[0:0x400] = fib
    word[0x400:fc_end] = text.encode("utf-16-le")
    word[0x800:0xA00] = chp_page
    word[0xA00:0xC00] = pap_page

    # The compound file around them
    END, FREE, FATSECT, NOSTREAM = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    header = bytearray(512)
    header[0:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    header[0x18:0x1A] = u16(0x3E)                            # minor version
    header[0x1A:0x1C] = u16(3)
    header[0x1C:0x1E] = u16(0xFFFE)                          # byte order
    header[0x1E:0x20] = u16(9)                               # 512-byte sectors
    header[0x20:0x22] = u16(6)                               # 64-byte mini sectors
    header[0x2C:0x30] = u32(1)                               # one FAT sector
    header[0x30:0x34] = u32(1)                               # directory at sector 1
    header[0x38:0x3C] = u32(4096)
    header[0x3C:0x40] = u32(END)
    header[0x44:0x48] = u32(END)
    header[0x4C:0x50] = u32(0)                               # the FAT is sector 0
    for i in range(1, 109):
        header[0x4C + i * 4:0x50 + i * 4] = u32(FREE)

    fat = [FREE] * 128
    fat[0] = FATSECT
    fat[1] = END
    for s in range(2, 9): fat[s] = s + 1
    fat[9] = END
    for s in range(10, 17): fat[s] = s + 1
    fat[17] = END

    def entry(name: str, kind: int, left, right, child, start, size):
        e = bytearray(128)
        raw = name.encode("utf-16-le") + b"\0\0"
        e[0:len(raw)] = raw
        e[64:66] = u16(len(raw))
        e[66] = kind
        e[67] = 1                                            # black
        e[68:72] = u32(left)
        e[72:76] = u32(right)
        e[76:80] = u32(child)
        e[116:120] = u32(start)
        e[120:124] = u32(size)
        return bytes(e)

    directory = (
        entry("Root Entry", 5, NOSTREAM, NOSTREAM, 1, END, 0) +
        entry("WordDocument", 2, NOSTREAM, 2, NOSTREAM, 2, 4096) +
        entry("1Table", 2, NOSTREAM, NOSTREAM, NOSTREAM, 10, 4096) +
        bytes(128)
    )
    out = bytes(header) + b"".join(u32(v) for v in fat) + directory + bytes(word) + bytes(table)
    assert len(out) == 512 * 19
    (OUT / "legacy.doc").write_bytes(out)
    written(OUT / "legacy.doc")


def prose() -> None:
    odt()
    rtf()
    doc()

def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    print(f"Writing fixtures into {OUT}")
    for step in (pdfs, wasm_decoded_images, docx, xlsx, pptx, texts, images, audio, zips, prose):
        step()
    total = sum(p.stat().st_size for p in OUT.iterdir() if p.is_file())
    count = sum(1 for p in OUT.iterdir() if p.is_file())
    print(f"\n{count} files, {total:,} bytes total")
    if total > 3 * 1024 * 1024:
        print("WARNING: fixtures exceed the 3 MB budget", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
